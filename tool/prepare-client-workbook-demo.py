"""Build repeat-safe SQL for the client workbook's supported demo records.

Requires openpyxl. This generates SQL; it never connects to a database.
"""
from datetime import date, datetime
from pathlib import Path
import json
import sys

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT.parent / "resource/sheets/Sistem_Manajemen_Pesanan_dan_Stok.xlsx"
OUTPUT = ROOT / "backend/supabase/seed_client_workbook.sql"
workbook = openpyxl.load_workbook(SOURCE, data_only=True)


def rows(sheet):
    return [list(row) for row in workbook[sheet].iter_rows(min_row=8, values_only=True) if row[0]]


def sql(value):
    if value is None:
        return "null"
    if isinstance(value, bool):
        return str(value).lower()
    if isinstance(value, (int, float)):
        return str(value)
    if isinstance(value, (date, datetime)):
        value = value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return "'" + str(value).replace("'", "''") + "'"


def item_id(code):
    return f"(select id from public.inventory_items where item_code={sql(code)})"


items = {r[0]: r for r in rows("Item")}
partners = {r[0]: r for r in rows("Mitra")}
bom = rows("BOM")
sales = rows("Pesanan")
purchases = rows("Pembelian")
movements = rows("Mutasi Stok")
subcontracts = {r[0]: r for r in rows("Subkontrak")}
categories = {"Produk Jadi": "finished_good", "Komponen": "component", "Bahan": "material", "Part Beli": "purchased_part", "Consumable": "consumable", "Subkontrak": "subcontract"}
procurement = {"Buat": "make", "Beli": "buy", "Subkontrak": "subcontract"}
statements = ["-- Generated from the client workbook by tool/prepare-client-workbook-demo.py.",
              "-- Adds demo rows once; reruns preserve existing or subsequently edited records.",
              "begin;", "select pg_advisory_xact_lock(hashtext('client-workbook-demo'));",
              "set local app.audit_source = 'client-workbook-demo';"]

for row in partners.values():
    if row[2] == "Pelanggan":
        # Workbook phone placeholders are not phone numbers; leave them blank.
        statements.append("insert into public.clients(name,email,postal_code,address) values (" + ",".join(sql(v) for v in [row[1], row[8], row[5], row[4]]) + ") on conflict (name) do nothing;")

for row in items.values():
    supplier = partners[row[5]][1] if row[5] else ""
    values = [row[0], row[1], categories[row[2]], procurement[row[3]], row[4], row[15], row[16], row[17], supplier, row[14], row[19] == "Ya"]
    statements.append("insert into public.inventory_items(item_code,item_name,category,procurement_type,unit,opening_qty,minimum_qty,target_qty,supplier_name,unit_cost,active) values (" + ",".join(map(sql, values)) + ") on conflict (item_code) do nothing;")

# The existing app gives each product a client. The component's own product
# master supports producing CP-001 separately instead of issuing its raw material
# again when FG-002 is assembled.
products = [("FG-001", "C-001"), ("FG-002", "C-002"), ("CP-001", "C-002")]
for code, client_code in products:
    row = items[code]
    tasks = [{"content": row[7] + ": " + row[9], "time": round(row[11] * 60, 3)}]
    tasks += [{"content": "", "time": ""} for _ in range(53)]
    values = [partners[client_code][1], row[1], row[6], row[13] or 0, json.dumps(tasks, ensure_ascii=False)]
    statements.append("insert into public.products(client_name,product_name,product_number,unit_price,tasks) values (" + ",".join(map(sql, values)) + ") on conflict (client_name,product_number) do nothing;")


def add_branch(code, product_expr, parent_expr="null", ancestors=()):
    if code in ancestors:
        raise ValueError(f"Cyclic workbook BOM: {ancestors} -> {code}")
    for edge in [r for r in bom if r[1] == code]:
        child = item_id(edge[3])
        branch = f"product_id={product_expr} and inventory_item_id={child} and parent_bom_no is not distinct from {parent_expr}"
        note = f"[Workbook demo] {edge[0]} | {edge[1]} -> {edge[3]}"
        statements.append(f"insert into public.product_materials(product_id,inventory_item_id,quantity_per_unit,parent_bom_no,notes) select {product_expr},{child},{sql(edge[5])},{parent_expr},{sql(note)} where not exists (select 1 from public.product_materials where {branch});")
        next_parent = f"(select bom_no from public.product_materials where {branch})"
        add_branch(edge[3], product_expr, next_parent, ancestors + (code,))


for code, client_code in products:
    product = f"(select id from public.products where client_name={sql(partners[client_code][1])} and product_number={sql(items[code][6])})"
    add_branch(code, product)

# Received purchases generate their own stock entries through the existing trigger.
# Do not also import their receipt rows from Mutasi Stok.
statements.append("do $demo$ declare v_purchase_id uuid; begin")
for row in purchases:
    status = "cancelled" if row[11] == "Batal" else "received" if row[12] == row[7] else "partial" if row[12] else "ordered"
    note = f"[Workbook demo] {row[0]} | {row[17]}"
    values = [sql(row[1]), sql(row[2]), sql(row[4]), item_id(row[5]), sql(row[7]), sql(row[12]), sql(row[8]), sql(row[10]), sql(status), sql(row[16]), sql(note)]
    # Skip before INSERT, since the purchase's BEFORE trigger posts receipts.
    # ON CONFLICT alone would still run that trigger for an existing row.
    statements.append("insert into public.purchases(purchase_no,purchase_date,supplier_name,item_id,ordered_qty,received_qty,unit_price,due_date,status,pic,notes) select " + ",".join(values) + f" where not exists (select 1 from public.purchases where purchase_no={sql(row[1])}) returning id into v_purchase_id;")
    receipt = next((m for m in movements if m[2] == "Terima PO" and m[3] == row[0]), None)
    if receipt is None:
        receipt = next((m for m in movements if m[2] == "Terima subkontrak" and subcontracts[m[3]][2] == row[0]), None)
    if receipt:
        statements.append(f"if v_purchase_id is not null then update public.stock_movements set movement_date={sql(receipt[1])}, notes={sql('[Workbook demo] ' + receipt[0] + ' | ' + receipt[9])} where reference_type='purchase' and reference_id=v_purchase_id::text; end if;")
statements.append("end $demo$;")

movement_types = {"Pakai bahan": "material_issue", "Hasil produksi": "production_output", "Kirim subkontrak": "subcontract_out", "Kirim pelanggan": "sale_shipment", "Koreksi masuk": "adjustment_in", "Koreksi keluar": "adjustment_out"}
manual_count = 0
for row in movements:
    if not row[11] or row[2] not in movement_types:
        continue
    values = [sql("WB-" + row[0]), sql(row[1]), sql(movement_types[row[2]]), "'manual'", sql(row[3]), item_id(row[4]), sql(row[6]), sql(row[11]), sql(row[8]), sql(f"[Workbook demo] {row[0]} | {row[9]}")]
    # Skip existing rows before the negative-stock BEFORE trigger runs.
    statements.append("insert into public.stock_movements(movement_no,movement_date,movement_type,reference_type,reference_id,item_id,quantity,delta,pic,notes) select " + ",".join(values) + f" where not exists (select 1 from public.stock_movements where movement_no={sql('WB-' + row[0])});")
    manual_count += 1

# Sales rows represent the outstanding balance. Historical partial shipments
# are already in the ledger. Preserve original requested/shipped quantities in
# the order note, since this app has no partial-shipment quantity fields.
for row in sales:
    cancelled = row[11] == "Batal"
    quantity = row[7] if cancelled else row[13]
    status = "Cancelled" if cancelled else "In preparation" if row[12] else "Order request"
    note = f"[Workbook demo] {row[0]} | Original ordered: {row[7]}; shipped: {row[12]}; outstanding: {row[13]}. Source status: {row[15]}. Displayed quantity is the outstanding balance. Historical production and shipments are in Movements."
    minutes = round(quantity * items[row[5]][11] * 60)
    values = [sql(row[2]), sql(row[10]), sql(row[4]), sql(row[1]), sql(row[6]), sql(quantity), sql(quantity * row[8]), sql(status), sql(minutes), sql(note)]
    statements.append("insert into public.orders(order_date,delivery_date,client,order_number,product_name,quantity,order_amount,progress,required_manhours,order_contact) select " + ",".join(values) + f" where not exists (select 1 from public.orders where client={sql(row[4])} and order_number={sql(row[1])});")

statements.append("commit;")
OUTPUT.write_text("\n\n".join(statements) + "\n", encoding="utf-8")
print(json.dumps({"output": str(OUTPUT), "items": len(items), "clients": 2, "products": len(products), "bom_rows": sum(s.startswith('insert into public.product_materials') for s in statements), "purchases": len(purchases), "manual_movements": manual_count, "orders": len(sales)}))
